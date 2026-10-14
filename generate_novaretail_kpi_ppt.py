"""
Generate a 5-slide NovaRetail executive KPI deck from Excel data and a PPTX template.

Install:
    pip install pandas openpyxl python-pptx matplotlib

Expected Excel sheets and columns:
1) Commercial
   Quarter, Region, Revenue_USD_M, Units_Sold_K, Returns_Rate_pct,
   Marketing_Spend_USD_M
2) Customer_Experience
   Quarter, Region, NPS, CSAT_pct, Avg_Delivery_Days, Support_Tickets_K
3) Operations
   Quarter, Region, Stockout_Rate_pct, Return_Processing_Days,
   Warehouse_Capacity_Util_pct, Fraud_Incidents

Run:
    python generate_novaretail_kpi_ppt.py \
        --data NovaRetail_Data.xlsx \
        --template NovaRetail_Q2_2026_Executive_Deck.pptx \
        --output NovaRetail_Q2_2026_Alternative_KPI_Deck.pptx
"""

import argparse
import os
import tempfile
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

# -----------------------------
# Theme
# -----------------------------
NAVY = "202A63"
DARK_NAVY = "141B49"
TEAL = "1AA6A6"
GREEN = "2C8C63"
RED = "C74747"
ORANGE = "E08B3E"
BLUE = "2F6DB3"
INK = "16213E"
GRAY = "667085"
LIGHT_BG = "F7F8FA"
GRID = "D9DEE8"
WHITE = "FFFFFF"
REGION_ORDER = ["North America", "Europe", "APAC", "Latin America"]


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def read_data(workbook):
    """Read and validate the three required Excel sheets."""
    commercial = pd.read_excel(workbook, sheet_name="Commercial", engine="openpyxl")
    customer = pd.read_excel(workbook, sheet_name="Customer_Experience", engine="openpyxl")
    operations = pd.read_excel(workbook, sheet_name="Operations", engine="openpyxl")

    required = {
        "Commercial": [
            "Quarter", "Region", "Revenue_USD_M", "Units_Sold_K",
            "Returns_Rate_pct", "Marketing_Spend_USD_M"
        ],
        "Customer_Experience": [
            "Quarter", "Region", "NPS", "CSAT_pct",
            "Avg_Delivery_Days", "Support_Tickets_K"
        ],
        "Operations": [
            "Quarter", "Region", "Stockout_Rate_pct",
            "Return_Processing_Days", "Warehouse_Capacity_Util_pct",
            "Fraud_Incidents"
        ],
    }
    for name, frame in [
        ("Commercial", commercial),
        ("Customer_Experience", customer),
        ("Operations", operations),
    ]:
        missing = [c for c in required[name] if c not in frame.columns]
        if missing:
            raise ValueError(f"Sheet '{name}' is missing columns: {missing}")
        frame["Quarter"] = frame["Quarter"].astype(str).str.strip()
        frame["Region"] = frame["Region"].astype(str).str.strip()

    return commercial, customer, operations


def quarter_key(q):
    year, quarter = q.split("-Q")
    return int(year), int(quarter)


def latest_and_previous(commercial):
    quarters = sorted(commercial["Quarter"].unique(), key=quarter_key)
    if len(quarters) < 2:
        raise ValueError("At least two quarters are required.")
    return quarters[-1], quarters[-2]


def remove_all_slides(prs):
    """Keep the template theme and dimensions, but remove its existing slides."""
    slide_ids = list(prs.slides._sldIdLst)
    for slide_id in slide_ids:
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        prs.slides._sldIdLst.remove(slide_id)


def blank_layout(prs):
    return prs.slide_layouts[6]


def set_background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_text(slide, text, x, y, w, h, size=12, color=INK,
             bold=False, align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Pt(0)
    frame.margin_top = frame.margin_bottom = Pt(0)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.65, 0.34, 8.4, 0.5, 19, INK, True, font="Aptos Display")
    if subtitle:
        add_text(slide, subtitle, 0.65, 0.84, 8.6, 0.25, 8.5, GRAY)


def add_footer(slide, slide_number, dark=False):
    color = "8FA1DB" if dark else "8A93A6"
    add_text(slide, "NovaRetail Quarterly Business Review", 0.55, 7.14, 3.2, 0.14, 5.5, color)
    add_text(slide, str(slide_number), 12.55, 7.14, 0.28, 0.14, 5.5, color, align=PP_ALIGN.RIGHT)


def add_icon_circle(slide, symbol, x, y):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.82), Inches(0.82)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(NAVY)
    shape.line.color.rgb = rgb(NAVY)
    add_text(slide, symbol, x + 0.05, y + 0.16, 0.72, 0.4, 18, WHITE, True, PP_ALIGN.CENTER)


def add_bullets(slide, bullets, x=9.2, y=1.35, w=3.45, h=4.7):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(0)
    tf.margin_right = Pt(0)
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(10)
        p.font.color.rgb = rgb(INK)
        p.space_after = Pt(10)
        p.text = "• " + p.text


def save_bar_chart(labels, series, colors, output, title=None, y_label=None,
                   horizontal=False, value_formats=None, ymax=None):
    """Create a clean chart image for insertion into PowerPoint."""
    plt.figure(figsize=(9.2, 5.2), dpi=180)
    ax = plt.gca()
    x = np.arange(len(labels))
    count = len(series)
    width = 0.72 / count

    for i, (name, values) in enumerate(series.items()):
        offset = (i - (count - 1) / 2) * width
        if horizontal:
            bars = ax.barh(x + offset, values, height=width, label=name, color=colors[i])
        else:
            bars = ax.bar(x + offset, values, width=width, label=name, color=colors[i])
        fmt = value_formats[i] if value_formats else "{:.0f}"
        for bar, value in zip(bars, values):
            label = fmt.format(value)
            if horizontal:
                ax.text(bar.get_width() + max(values) * 0.012,
                        bar.get_y() + bar.get_height() / 2, label,
                        va="center", fontsize=8)
            else:
                ax.text(bar.get_x() + bar.get_width() / 2,
                        bar.get_height() + max(values) * 0.018, label,
                        ha="center", va="bottom", fontsize=8)

    if horizontal:
        ax.set_yticks(x, labels)
        ax.invert_yaxis()
    else:
        ax.set_xticks(x, labels)
    if ymax:
        (ax.set_xlim if horizontal else ax.set_ylim)(0, ymax)
    if title:
        ax.set_title(title, fontsize=11, fontweight="bold")
    if y_label:
        ax.set_ylabel(y_label, fontsize=9)
    if count > 1:
        ax.legend(frameon=False, loc="upper left", fontsize=8)
    ax.grid(axis="x" if horizontal else "y", color="#D9DEE8", linewidth=0.6)
    ax.set_axisbelow(True)
    ax.spines[["top", "right", "left"]].set_visible(False)
    if not horizontal:
        ax.spines["bottom"].set_color("#B8C0CC")
    plt.tight_layout()
    plt.savefig(output, transparent=True, bbox_inches="tight")
    plt.close()


def q2_rows(frame, quarter):
    rows = frame[frame["Quarter"] == quarter].copy()
    return rows.set_index("Region").reindex(REGION_ORDER).reset_index()


def generate_deck(data_path, template_path, output_path):
    commercial, customer, operations = read_data(data_path)
    latest, previous = latest_and_previous(commercial)

    c_latest = q2_rows(commercial, latest)
    c_prev = q2_rows(commercial, previous)
    cx_latest = q2_rows(customer, latest)
    op_latest = q2_rows(operations, latest)

    total_units = c_latest["Units_Sold_K"].sum()
    prev_units = c_prev["Units_Sold_K"].sum()
    unit_growth = (total_units / prev_units - 1) * 100
    total_revenue = c_latest["Revenue_USD_M"].sum()
    total_marketing = c_latest["Marketing_Spend_USD_M"].sum()
    marketing_productivity = total_revenue / total_marketing
    fraud_total = op_latest["Fraud_Incidents"].sum()

    revenue_per_marketing = (
        c_latest["Revenue_USD_M"] / c_latest["Marketing_Spend_USD_M"]
    )
    ticket_total = cx_latest["Support_Tickets_K"].sum()
    latam_tickets = cx_latest.loc[cx_latest["Region"] == "Latin America", "Support_Tickets_K"].iloc[0]
    latam_ticket_share = latam_tickets / ticket_total * 100

    prs = Presentation(template_path)
    remove_all_slides(prs)
    layout = blank_layout(prs)

    with tempfile.TemporaryDirectory() as temp_dir:
        temp_dir = Path(temp_dir)

        # Slide 1: Executive summary
        slide = prs.slides.add_slide(layout)
        set_background(slide, NAVY)
        add_text(slide, "NovaRetail — Quarterly Business Review", 0.6, 0.55, 7.0, 0.45,
                 23, WHITE, True, font="Aptos Display")
        add_text(slide, f"{latest} • Executive KPI Review", 0.6, 1.02, 4.0, 0.22,
                 8, "A8B2D8")

        cards = [
            (f"{total_units / 1000:.2f}M", f"{latest} Units Sold"),
            (f"${marketing_productivity:.2f}", "Revenue per $1 Marketing"),
            (f"{fraud_total:.0f}", "Fraud Incidents"),
        ]
        for i, (value, label) in enumerate(cards):
            x = 0.62 + i * 3.72
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.52),
                Inches(3.35), Inches(1.02)
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(DARK_NAVY)
            shape.line.color.rgb = rgb(DARK_NAVY)
            add_text(slide, value, x + 0.12, 1.72, 3.1, 0.34, 19, TEAL, True, PP_ALIGN.CENTER)
            add_text(slide, label, x + 0.12, 2.17, 3.1, 0.18, 7, "B9C2E1", align=PP_ALIGN.CENTER)

        add_text(slide, "EXECUTIVE SUMMARY", 0.62, 2.88, 2.2, 0.18, 6.5, "8FA1DB", True)
        apac_capacity = op_latest.loc[op_latest["Region"] == "APAC", "Warehouse_Capacity_Util_pct"].iloc[0]
        latam_csat = cx_latest.loc[cx_latest["Region"] == "Latin America", "CSAT_pct"].iloc[0]
        latam_processing = op_latest.loc[op_latest["Region"] == "Latin America", "Return_Processing_Days"].iloc[0]
        latam_fraud = op_latest.loc[op_latest["Region"] == "Latin America", "Fraud_Incidents"].iloc[0]
        summary = (
            f"{latest} volume reached {total_units / 1000:.2f}M units, up {unit_growth:.1f}% quarter over quarter. "
            f"Marketing productivity was ${marketing_productivity:.2f} revenue per $1 spent. "
            f"APAC combines strong demand with {apac_capacity:.0f}% warehouse utilization, while Latin America "
            f"combines {latam_csat:.0f}% CSAT, {latam_processing:.0f}-day return processing and "
            f"{latam_fraud:.0f} fraud incidents."
        )
        add_text(slide, summary, 0.62, 3.25, 9.35, 1.65, 12, WHITE, valign=MSO_ANCHOR.MIDDLE)
        add_icon_circle(slide, "↗", 10.75, 3.52)
        add_footer(slide, 1, dark=True)

        # Slide 2: Demand and marketing productivity
        chart1 = temp_dir / "marketing_productivity.png"
        save_bar_chart(
            REGION_ORDER,
            {"Revenue per $1 marketing": revenue_per_marketing.tolist()},
            ["#1AA6A6"], chart1,
            y_label="Revenue / marketing spend",
            value_formats=["{:.1f}x"], ymax=max(revenue_per_marketing) * 1.25
        )
        slide = prs.slides.add_slide(layout)
        set_background(slide, LIGHT_BG)
        add_title(slide, "Demand Growth Improves, but Marketing Productivity Splits",
                  f"{latest} regional commercial efficiency")
        slide.shapes.add_picture(str(chart1), Inches(0.68), Inches(1.24), width=Inches(8.1))
        units_map = c_latest.set_index("Region")["Units_Sold_K"]
        apac_prev = c_prev.set_index("Region").loc["APAC", "Units_Sold_K"]
        apac_latest = units_map["APAC"]
        latam_prev = c_prev.set_index("Region").loc["Latin America", "Units_Sold_K"]
        latam_latest = units_map["Latin America"]
        bullets = [
            f"North America sold {units_map['North America'] / 1000:.2f}M units and remained the largest volume market.",
            f"Europe delivered ${revenue_per_marketing.iloc[1]:.1f} revenue per $1 of marketing spend.",
            f"APAC reached {apac_latest:.0f}K units, up {(apac_latest / apac_prev - 1) * 100:.1f}% quarter over quarter.",
            f"Latin America volume declined {(1 - latam_latest / latam_prev) * 100:.1f}% quarter over quarter."
        ]
        add_bullets(slide, bullets)
        add_icon_circle(slide, "$", 10.55, 5.65)
        add_footer(slide, 2)

        # Slide 3: Customer service
        chart2a = temp_dir / "csat.png"
        chart2b = temp_dir / "tickets.png"
        save_bar_chart(REGION_ORDER, {"CSAT": cx_latest["CSAT_pct"].tolist()},
                       ["#2C8C63"], chart2a, title="CSAT (%)",
                       value_formats=["{:.0f}%"], ymax=105)
        save_bar_chart(REGION_ORDER, {"Tickets": cx_latest["Support_Tickets_K"].tolist()},
                       ["#E08B3E"], chart2b, title="Support tickets (K)",
                       value_formats=["{:.0f}"], ymax=35)
        slide = prs.slides.add_slide(layout)
        set_background(slide, LIGHT_BG)
        add_title(slide, "Service Performance Diverges Across Regions",
                  f"{latest} customer experience and service demand")
        slide.shapes.add_picture(str(chart2a), Inches(0.62), Inches(1.25), width=Inches(4.25))
        slide.shapes.add_picture(str(chart2b), Inches(4.55), Inches(1.25), width=Inches(4.25))
        cx = cx_latest.set_index("Region")
        bullets = [
            f"APAC leads at {cx.loc['APAC', 'CSAT_pct']:.0f}% CSAT with {cx.loc['APAC', 'Avg_Delivery_Days']:.1f}-day delivery.",
            f"North America and Europe remain at {cx.loc['North America', 'CSAT_pct']:.0f}% and {cx.loc['Europe', 'CSAT_pct']:.0f}% CSAT.",
            f"Latin America fell to {cx.loc['Latin America', 'CSAT_pct']:.0f}% CSAT as delivery reached {cx.loc['Latin America', 'Avg_Delivery_Days']:.1f} days.",
            f"Latin America generated {latam_tickets:.0f}K tickets, {latam_ticket_share:.0f}% of the regional total."
        ]
        add_bullets(slide, bullets)
        add_icon_circle(slide, "★", 10.55, 5.65)
        add_footer(slide, 3)

        # Slide 4: Operations and risk
        chart3 = temp_dir / "operations_risk.png"
        save_bar_chart(
            REGION_ORDER,
            {
                "Return processing days": op_latest["Return_Processing_Days"].tolist(),
                "Fraud incidents": op_latest["Fraud_Incidents"].tolist(),
            },
            ["#2F6DB3", "#C74747"], chart3,
            value_formats=["{:.0f}", "{:.0f}"], ymax=35
        )
        slide = prs.slides.add_slide(layout)
        set_background(slide, LIGHT_BG)
        add_title(slide, "Returns Processing and Fraud Define the Risk Gap",
                  f"{latest} operational control indicators")
        slide.shapes.add_picture(str(chart3), Inches(0.68), Inches(1.25), width=Inches(8.1))
        op = op_latest.set_index("Region")
        latam_fraud_start = operations[
            operations["Region"] == "Latin America"
        ].sort_values("Quarter", key=lambda x: x.map(quarter_key)).iloc[0]["Fraud_Incidents"]
        bullets = [
            f"Latin America has the slowest returns cycle at {op.loc['Latin America', 'Return_Processing_Days']:.0f} days and {op.loc['Latin America', 'Fraud_Incidents']:.0f} fraud incidents.",
            f"Latin America fraud incidents increased from {latam_fraud_start:.0f} to {op.loc['Latin America', 'Fraud_Incidents']:.0f} across the dataset.",
            f"APAC return processing is {op.loc['APAC', 'Return_Processing_Days']:.0f} days, but warehouse utilization reached {op.loc['APAC', 'Warehouse_Capacity_Util_pct']:.0f}%.",
            f"North America and Europe recorded {op.loc['North America', 'Fraud_Incidents']:.0f} and {op.loc['Europe', 'Fraud_Incidents']:.0f} fraud incidents."
        ]
        add_bullets(slide, bullets)
        add_icon_circle(slide, "!", 10.55, 5.65)
        add_footer(slide, 4)

        # Slide 5: Recommendations
        slide = prs.slides.add_slide(layout)
        set_background(slide, NAVY)
        add_text(slide, "Recommended Actions & Next Steps", 0.65, 0.55, 6.5, 0.45,
                 22, WHITE, True, font="Aptos Display")
        actions = [
            ("PROTECT APAC CAPACITY", "Add overflow capacity and define a durable expansion trigger."),
            ("RECOVER LATIN AMERICA SERVICE", "Reduce delivery time, support tickets and return-processing days."),
            ("STRENGTHEN FRAUD CONTROLS", "Prioritize monitoring and exception review in Latin America."),
            ("REBALANCE MARKETING", "Shift incremental spend toward regions with stronger revenue-per-dollar productivity."),
        ]
        y = 1.45
        for heading, detail in actions:
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.72), Inches(y + 0.04), Inches(0.23), Inches(0.23))
            dot.fill.solid()
            dot.fill.fore_color.rgb = rgb(TEAL)
            dot.line.color.rgb = rgb(TEAL)
            add_text(slide, f"{heading}  •  {detail}", 1.08, y, 10.6, 0.55, 11, WHITE)
            y += 1.0
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(5.85), Inches(10.6), Inches(0.01))
        line.fill.solid()
        line.fill.fore_color.rgb = rgb("6C78AC")
        line.line.color.rgb = rgb("6C78AC")
        add_text(slide,
                 "Sustain commercial momentum while closing the service and control gap in Latin America.",
                 0.72, 6.05, 10.7, 0.55, 13, WHITE, True)
        add_icon_circle(slide, "✓", 11.45, 2.7)
        add_footer(slide, 5, dark=True)

        prs.save(output_path)

    print(f"Presentation created: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="Generate NovaRetail executive KPI PowerPoint")
    parser.add_argument("--data", required=True, help="Excel workbook containing the three data sheets")
    parser.add_argument("--template", required=True, help="Existing PPTX template")
    parser.add_argument("--output", default="NovaRetail_KPI_Deck.pptx", help="Output PPTX path")
    args = parser.parse_args()

    for path in [args.data, args.template]:
        if not os.path.exists(path):
            raise FileNotFoundError(path)
    generate_deck(args.data, args.template, args.output)


if __name__ == "__main__":
    main()
